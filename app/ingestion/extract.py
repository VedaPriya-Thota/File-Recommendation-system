"""
Content extraction boundary (Phase 10). Pure: bytes in, ExtractedContent out;
no filesystem, database or network. Supported: .txt .md .pdf .docx (.pptx and
others are deliberately out of scope -- they report "unsupported").

Failures are reported as a short `error_code` only -- never the exception
message, which could contain paths or document content.
"""
import io
from dataclasses import dataclass
from typing import Optional

from app.ingestion.normalize import normalize_text

SUPPORTED_EXTENSIONS = ("txt", "md", "pdf", "docx")

OK, EMPTY, UNSUPPORTED, FAILED = "ok", "empty", "unsupported", "failed"


@dataclass(frozen=True)
class ExtractedContent:
    status: str                      # ok | empty | unsupported | failed
    text: str = ""                   # NORMALIZED text (empty unless status == ok)
    error_code: Optional[str] = None  # set only when status == failed


def _extract_plain(data: bytes) -> str:
    return data.decode("utf-8-sig", errors="replace")


def _extract_pdf(data: bytes) -> str:
    from pypdf import PdfReader  # lazy: a missing optional library degrades to failed("extractor_unavailable")

    reader = PdfReader(io.BytesIO(data))
    if reader.is_encrypted:
        if not reader.decrypt(""):
            raise PermissionError("encrypted_pdf")
    return "\n\n".join((page.extract_text() or "") for page in reader.pages)


def _extract_docx(data: bytes) -> str:
    from docx import Document

    doc = Document(io.BytesIO(data))
    parts = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            parts.append(" ".join(cell.text for cell in row.cells))
    return "\n".join(parts)


_EXTRACTORS = {"txt": _extract_plain, "md": _extract_plain, "pdf": _extract_pdf, "docx": _extract_docx}


class ContentExtractor:
    def supports(self, extension: str) -> bool:
        return (extension or "").lower() in _EXTRACTORS

    def extract(self, extension: str, data: bytes) -> ExtractedContent:
        ext = (extension or "").lower()
        fn = _EXTRACTORS.get(ext)
        if fn is None:
            return ExtractedContent(UNSUPPORTED)
        try:
            raw = fn(data)
        except ImportError:
            return ExtractedContent(FAILED, error_code="extractor_unavailable")
        except PermissionError:
            return ExtractedContent(FAILED, error_code="encrypted_" + ext)
        except Exception:  # noqa: BLE001 -- any parser error on a bad file is a per-file failure, never a crash
            return ExtractedContent(FAILED, error_code=f"malformed_{ext}")
        text = normalize_text(raw)
        return ExtractedContent(OK, text) if text else ExtractedContent(EMPTY)
